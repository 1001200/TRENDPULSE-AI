import io
import csv
from datetime import datetime
from sqlalchemy.orm import Session
from backend.app.db.session import SessionLocal, engine, Base
from backend.app.models.models import Product, SaleOrder, InventoryItem, DataQualityLog
from backend.app.pipeline.extractors import (
    detect_file_format,
    extract_csv,
    extract_excel,
    extract_pdf,
    extract_docx,
    extract_txt,
    extract_image,
    detect_record_type
)
from backend.app.pipeline.staging import (
    analyze_uploaded_files,
    update_staged_batch,
    commit_staged_batch
)

def run_tests():
    print("=== STARTING TRENDPULSE IMPORTER PIPELINE TESTS ===")
    db = SessionLocal()

    # 1. Test CSV Extraction & Column Sniffing
    print("\n[1] Testing CSV Extraction...")
    csv_data = "Item Title, Garment Category, Sizing, Selling Price, Units Sold, Sale Date, Customer Handle\n" \
               "vintage oversized hoodie, outerwear, xl, $1,450.00, 2, 2026-10-15, @maya_vintage\n" \
               "y2k graphic baby tee, tops, Medium, 890, 1, 15/10/2026, @alina_resale\n" \
               "y2k graphic baby tee, tops, Medium, 890, 1, 15/10/2026, @alina_resale\n" # duplicate row
    headers, rows, warn = extract_csv(csv_data.encode("utf-8"))
    assert len(headers) == 7, f"Expected 7 headers, got {len(headers)}"
    assert len(rows) == 3, f"Expected 3 rows, got {len(rows)}"
    print(f"[OK] CSV extracted: {len(headers)} columns, {len(rows)} rows.")

    # 2. Test Excel (.xlsx) Extraction
    print("\n[2] Testing XLSX Extraction...")
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Product Name", "Category", "Size", "Price", "Quantity", "Order Date", "Buyer"])
    ws.append(["Vintage Levi 501 Jeans", "Bottoms", "32", 2400, 1, "2026-10-18", "@denim_head"])
    ws.append(["90s Leather Jacket", "Outerwear", "L", 4200, 1, "2026-10-19", "@leather_fan"])
    xlsx_buf = io.BytesIO()
    wb.save(xlsx_buf)
    xlsx_bytes = xlsx_buf.getvalue()

    headers_x, rows_x, _ = extract_excel(xlsx_bytes)
    assert len(rows_x) == 2, f"Expected 2 rows, got {len(rows_x)}"
    print(f"[OK] XLSX extracted: {len(headers_x)} columns, {len(rows_x)} rows.")

    # 3. Test Word Document (.docx) Extraction
    print("\n[3] Testing DOCX Extraction...")
    import docx
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Item"
    hdr_cells[1].text = "Size"
    hdr_cells[2].text = "Retail Price"
    hdr_cells[3].text = "Channel"
    row_cells = table.add_row().cells
    row_cells[0].text = "Floral Prairie Midi Dress"
    row_cells[1].text = "Small"
    row_cells[2].text = "1850"
    row_cells[3].text = "Story Drop"
    docx_buf = io.BytesIO()
    doc.save(docx_buf)
    docx_bytes = docx_buf.getvalue()

    headers_d, rows_d, _, _ = extract_docx(docx_bytes)
    assert len(rows_d) == 1, f"Expected 1 row in docx, got {len(rows_d)}"
    print(f"[OK] DOCX extracted: {len(headers_d)} columns, {len(rows_d)} rows.")

    # 4. Test PDF Extraction
    print("\n[4] Testing PDF Extraction...")
    from pypdf import PdfWriter
    pdf_writer = PdfWriter()
    pdf_writer.add_blank_page(width=72, height=72)
    pdf_buf = io.BytesIO()
    pdf_writer.write(pdf_buf)
    pdf_bytes = pdf_buf.getvalue()
    # verify format detection on PDF
    fmt = detect_file_format("invoice.pdf", pdf_bytes)
    assert fmt == "pdf", f"Expected pdf, got {fmt}"
    print(f"[OK] PDF format verification OK.")

    # 5. Test Image Extraction & Graceful Non-Business Handling
    print("\n[5] Testing Image Extraction & Safety...")
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (200, 100), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((10, 10), "Vintage Denim Jacket 1200", fill=(0, 0, 0))
    img_buf = io.BytesIO()
    img.save(img_buf, format="PNG")
    img_bytes = img_buf.getvalue()

    fmt_img = detect_file_format("receipt.png", img_bytes)
    assert fmt_img == "png"
    h_img, r_img, warn_img, _ = extract_image(img_bytes, "receipt.png")
    print(f"[OK] Image extraction executed. Warning/Result: {warn_img or 'Parsed'}")

    # Test non-business blank image
    blank_img = Image.new("RGB", (50, 50), color=(128, 128, 128))
    blank_buf = io.BytesIO()
    blank_img.save(blank_buf, format="JPEG")
    blank_bytes = blank_buf.getvalue()
    h_b, r_b, warn_b, _ = extract_image(blank_bytes, "model_photo.jpg")
    assert len(r_b) == 0, "Expected 0 records for non-business image"
    assert "does not contain" in warn_b or "No" in warn_b
    print("[OK] Non-business photo safely rejected without fictional data generation.")

    # 6. Test Multi-File Staging & Analysis
    print("\n[6] Testing Multi-File Staging & Quality Analysis...")
    files_payload = [
        ("batch1_sales.csv", csv_data.encode("utf-8")),
        ("batch2_inventory.xlsx", xlsx_bytes)
    ]
    staged = analyze_uploaded_files(files_payload, db)
    assert staged["success"] is True
    session_id = staged["session_id"]
    summary = staged["summary"]
    issues = staged["issues"]
    print(f"[OK] Staged session created: {session_id}")
    print(f"  Total records detected: {summary['total_rows']}")
    print(f"  Valid records ready: {summary['valid_rows']}")
    print(f"  Confirmed duplicates: {summary['confirmed_duplicates']}")
    print(f"  Formatting cleanings: {summary['formatting_cleaned']}")
    print(f"  Quality Score: {summary['quality_score']}%")

    # 7. Test Staging Updates (Accept Correction & Exclude Duplicates)
    print("\n[7] Testing Staging Corrections & Exclusions...")
    update_res = update_staged_batch(session_id, {
        "bulk_action": "exclude_confirmed_duplicates"
    })
    assert update_res["success"] is True
    print("[OK] Excluded confirmed duplicates via bulk action.")

    # 8. Test Atomic Database Commit
    print("\n[8] Testing Atomic Database Commit...")
    commit_res = commit_staged_batch(session_id, db)
    assert commit_res["success"] is True
    print(f"[OK] Database commit succeeded: {commit_res['imported_count']} records imported.")
    print(f"  Saved Products: {commit_res['entities_saved']['products']}")
    print(f"  Saved Sales: {commit_res['entities_saved']['sales']}")
    print(f"  Saved Inventory: {commit_res['entities_saved']['inventory']}")

    # Check DataQualityLog audit trail
    log = db.query(DataQualityLog).filter(DataQualityLog.batch_id == commit_res["batch_id"]).first()
    assert log is not None
    assert log.valid_rows == commit_res["imported_count"]
    print(f"[OK] Audit trail recorded in DataQualityLog with {len(log.audit_trail_json)} mutations.")

    print("\n=== ALL PIPELINE TESTS PASSED SUCCESSFULLY! ===")
    db.close()

if __name__ == "__main__":
    run_tests()
