import io
import re
import csv
import logging
from typing import List, Dict, Any, Tuple, Optional

logger = logging.getLogger(__name__)

# Maximum file size permitted: 25 MB
MAX_FILE_SIZE_BYTES = 25 * 1024 * 1024

SUPPORTED_EXTENSIONS = {
    ".csv", ".xlsx", ".xls", ".pdf", ".docx", ".txt",
    ".png", ".jpg", ".jpeg", ".webp"
}


def detect_file_format(filename: str, content: bytes) -> str:
    """
    Validates file extension and magic byte signatures to prevent malicious/misnamed files.
    """
    fn_lower = filename.lower()
    ext = "." + fn_lower.rsplit(".", 1)[-1] if "." in fn_lower else ""

    if not content:
        raise ValueError(f"Uploaded file '{filename}' is completely empty (0 bytes).")

    if len(content) > MAX_FILE_SIZE_BYTES:
        raise ValueError(f"File size exceeds the maximum allowed limit of 25MB.")

    # Magic byte inspection
    if content.startswith(b"%PDF-"):
        return "pdf"
    if content.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if content.startswith(b"\xff\xd8\xff"):
        return "jpg"
    if len(content) > 12 and content[:4] == b"RIFF" and content[8:12] == b"WEBP":
        return "webp"
    if content.startswith(b"PK\x03\x04"):
        # Could be xlsx or docx
        if ext in (".xlsx", ".xlsm"):
            return "xlsx"
        if ext in (".docx", ".docm"):
            return "docx"
        return "xlsx" if "sheet" in ext or "xl" in ext else "docx"
    if content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "xls"

    # Fallback to extension matching for text/csv
    if ext == ".csv":
        return "csv"
    if ext in (".txt", ".log", ".tsv"):
        return "txt"
    if ext in (".xlsx", ".xls"):
        return ext.lstrip(".")
    if ext in (".png", ".jpg", ".jpeg", ".webp"):
        return ext.lstrip(".")
    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"

    # Check if decodable text
    try:
        content[:1024].decode("utf-8")
        return "csv" if "," in content[:1024].decode("utf-8") else "txt"
    except UnicodeDecodeError:
        pass

    raise ValueError(f"Unsupported file format for '{filename}'. Allowed: CSV, XLS, XLSX, PDF, DOCX, TXT, PNG, JPG, JPEG, WebP")


def _decode_text_stream(content: bytes) -> str:
    """Safely decodes raw bytes into string trying common retail file encodings."""
    encodings = ["utf-8-sig", "utf-8", "latin-1", "cp1252", "iso-8859-1"]
    for enc in encodings:
        try:
            return content.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return content.decode("utf-8", errors="replace")


def extract_csv(content: bytes) -> Tuple[List[str], List[Dict[str, Any]], Optional[str]]:
    """Extracts tabular records from CSV with delimiter sniffer."""
    text = _decode_text_stream(content)
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return [], [], "File contains no non-empty lines."

    # Sniff delimiter
    delimiter = ","
    sample = "\n".join(lines[:10])
    try:
        sniffer = csv.Sniffer()
        dialect = sniffer.sniff(sample, delimiters=[",", "\t", ";", "|"])
        delimiter = dialect.delimiter
    except Exception:
        # Fallback check
        first_line = lines[0]
        if "\t" in first_line and first_line.count("\t") > first_line.count(","):
            delimiter = "\t"
        elif ";" in first_line and first_line.count(";") > first_line.count(","):
            delimiter = ";"
        elif "|" in first_line and first_line.count("|") > first_line.count(","):
            delimiter = "|"

    reader = csv.reader(lines, delimiter=delimiter)
    all_rows = list(reader)
    if not all_rows:
        return [], [], "No tabular rows found in CSV."

    raw_headers = all_rows[0]
    headers = [str(h).strip() if str(h).strip() else f"Column_{i+1}" for i, h in enumerate(raw_headers)]

    rows = []
    for row_idx, r in enumerate(all_rows[1:], start=2):
        if not any(str(c).strip() for c in r):
            continue
        row_dict = {}
        for i, h in enumerate(headers):
            row_dict[h] = r[i].strip() if i < len(r) else ""
        rows.append(row_dict)

    warning = None
    if not rows:
        warning = "CSV has header columns but no data rows."
    return headers, rows, warning


def extract_excel(content: bytes, is_xls: bool = False) -> Tuple[List[str], List[Dict[str, Any]], Optional[str]]:
    """Extracts tabular rows from XLSX or XLS workbooks."""
    rows = []
    headers = []
    warning = None

    if is_xls:
        # Old Excel 97-2004 format (.xls)
        try:
            import xlrd
            wb = xlrd.open_workbook(file_contents=content)
            ws = wb.sheet_by_index(0)
            if ws.nrows == 0:
                return [], [], "Excel sheet contains no data."

            raw_headers = [str(ws.cell_value(0, col)).strip() for col in range(ws.ncols)]
            headers = [h if h else f"Column_{i+1}" for i, h in enumerate(raw_headers)]

            for row_idx in range(1, ws.nrows):
                row_vals = [ws.cell_value(row_idx, col) for col in range(ws.ncols)]
                if not any(v not in (None, "") for v in row_vals):
                    continue
                row_dict = {}
                for i, h in enumerate(headers):
                    val = row_vals[i] if i < len(row_vals) else None
                    row_dict[h] = val
                rows.append(row_dict)
            return headers, rows, warning
        except Exception as e:
            logger.warning(f"xlrd failed to parse XLS: {e}")
            # Try openpyxl as fallback

    # XLSX / OpenPyXL
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        iter_rows = ws.iter_rows(values_only=True)
        try:
            first_row = next(iter_rows)
        except StopIteration:
            return [], [], "Excel file is empty."

        headers = [str(cell).strip() if cell is not None and str(cell).strip() else f"Column_{i+1}" for i, cell in enumerate(first_row)]

        for r in iter_rows:
            if not any(v is not None and str(v).strip() != "" for v in r):
                continue
            row_dict = {}
            for i, h in enumerate(headers):
                val = r[i] if i < len(r) else None
                row_dict[h] = val
            rows.append(row_dict)
        wb.close()
    except Exception as e:
        raise ValueError(f"Failed to read Excel file: {str(e)}")

    return headers, rows, warning


def extract_pdf(content: bytes) -> Tuple[List[str], List[Dict[str, Any]], Optional[str], str]:
    """
    Extracts text and tables from digital PDF documents using pypdf.
    Handles password-protected, empty, and corrupt PDFs gracefully.
    Supports multi-page structured tables via layout extraction mode.
    """
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise ValueError("This PDF document is password-protected or encrypted. Please remove the password protection and re-upload.")

        total_pages = len(reader.pages)
        if total_pages == 0:
            return [], [], "PDF document has no pages.", ""

        # Step 1: Layout extraction across pages
        pages_layout_lines = []
        all_layout_text = []
        for p_idx, page in enumerate(reader.pages):
            try:
                txt = page.extract_text(extraction_mode="layout") or ""
            except Exception:
                txt = page.extract_text() or ""
            all_layout_text.append(txt)
            pages_layout_lines.append([l.strip() for l in txt.splitlines() if l.strip()])

        full_text = "\n".join(all_layout_text).strip()
        if not full_text:
            return [], [], "PDF appears to be scanned or contains no selectable digital text. OCR extraction required.", ""

        # Step 2: Try detecting multi-column table headers across pages
        retail_header_keywords = {
            'order', 'order id', 'product', 'item', 'date', 'category', 'size',
            'qty', 'quantity', 'price', 'unit price', 'cost', 'unit cost',
            'sku', 'status', 'payment', 'sales value', 'amount', 'total'
        }

        candidate_headers = []
        for pl in pages_layout_lines:
            for line in pl:
                parts = re.split(r'\s{2,}', line)
                if len(parts) >= 3:
                    parts_lower = [p.lower() for p in parts]
                    matches = sum(1 for p in parts_lower if any(kw in p for kw in retail_header_keywords))
                    if matches >= 3:
                        candidate_headers.append((tuple(parts), matches, len(parts)))

        if candidate_headers:
            from collections import Counter
            header_tuples = [h[0] for h in candidate_headers]
            best_header_tuple, _ = Counter(header_tuples).most_common(1)[0]
            headers = [h.strip() or f"Column_{i+1}" for i, h in enumerate(best_header_tuple)]
            col_count = len(headers)

            rows = []
            header_set = set(header_tuples)
            for pl in pages_layout_lines:
                for line in pl:
                    parts = re.split(r'\s{2,}', line)
                    if tuple(parts) in header_set:
                        continue
                    if len(parts) == col_count:
                        row_dict = {headers[i]: parts[i].strip() for i in range(col_count)}
                        rows.append(row_dict)
                    elif len(parts) > col_count and len(parts) >= 4:
                        row_dict = {headers[i]: (parts[i].strip() if i < len(parts) else "") for i in range(col_count)}
                        rows.append(row_dict)

            if rows:
                return headers, rows, None, full_text

        # Step 3: Fallback to _parse_text_lines_to_table
        headers, rows = _parse_text_lines_to_table(full_text)
        warning = None
        if not rows:
            warning = "Extracted digital text from PDF, but no consistent tabular business records were detected."
        return headers, rows, warning, full_text
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Could not read PDF document (corrupt or unreadable format): {str(e)}")


def extract_docx(content: bytes) -> Tuple[List[str], List[Dict[str, Any]], Optional[str], str]:
    """
    Extracts structured business records and tables from Word documents (.docx).
    """
    try:
        import docx
        doc = docx.Document(io.BytesIO(content))
        headers = []
        rows = []
        full_text = []

        # 1. First check Word Tables
        if doc.tables:
            for table in doc.tables:
                if len(table.rows) >= 2:
                    t_headers = [cell.text.strip() or f"Column_{i+1}" for i, cell in enumerate(table.rows[0].cells)]
                    for r_idx in range(1, len(table.rows)):
                        cell_vals = [cell.text.strip() for cell in table.rows[r_idx].cells]
                        if not any(cell_vals):
                            continue
                        row_dict = {}
                        for i, h in enumerate(t_headers):
                            row_dict[h] = cell_vals[i] if i < len(cell_vals) else ""
                        rows.append(row_dict)
                    if t_headers and not headers:
                        headers = t_headers

        # Collect paragraph text
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text.strip())

        text_content = "\n".join(full_text)

        # 2. If no tables found, try parsing text paragraphs
        if not rows and text_content:
            p_headers, p_rows = _parse_text_lines_to_table(text_content)
            if p_rows:
                headers, rows = p_headers, p_rows

        warning = None
        if not rows:
            warning = "Document text was extracted, but could not be mapped reliably to structured business fields."

        return headers, rows, warning, text_content
    except Exception as e:
        raise ValueError(f"Could not read Word document (.docx): {str(e)}")


def extract_txt(content: bytes) -> Tuple[List[str], List[Dict[str, Any]], Optional[str]]:
    """Extracts records from plain text files (.txt, .tsv, etc.)."""
    text = _decode_text_stream(content)
    if not text.strip():
        return [], [], "Text file is empty."

    # Try standard delimiter parsing first (like CSV / TSV / pipe)
    headers, rows, _ = extract_csv(content)
    if rows and len(headers) > 1:
        return headers, rows, None

    # Fallback to key-value or pattern line parser
    h, r = _parse_text_lines_to_table(text)
    warning = None if r else "Text file could not be parsed into recognizable tabular business records."
    return h, r, warning


def extract_image(content: bytes, filename: str) -> Tuple[List[str], List[Dict[str, Any]], Optional[str], str]:
    """
    Performs OCR on receipts, screenshots, invoices, and photographed tables.
    If the image is a product photo or contains no business data, returns clear notification.
    """
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(content))
        img.verify()
    except Exception as e:
        raise ValueError(f"Image file is corrupt or unreadable: {str(e)}")

    extracted_text = ""
    ocr_used = None

    # 1. Try RapidOCR (pure onnx standalone OCR)
    try:
        from rapidocr_onnxruntime import RapidOCR
        engine = RapidOCR()
        ocr_result, _ = engine(content)
        if ocr_result:
            extracted_text = "\n".join([line[1] for line in ocr_result])
            ocr_used = "RapidOCR (ONNX)"
    except Exception as e:
        logger.debug(f"RapidOCR unavailable: {e}")

    # 2. Try pytesseract if RapidOCR wasn't used
    if not extracted_text:
        try:
            import pytesseract
            img = Image.open(io.BytesIO(content))
            extracted_text = pytesseract.image_to_string(img)
            if extracted_text.strip():
                ocr_used = "Tesseract OCR"
        except Exception as e:
            logger.debug(f"pytesseract unavailable: {e}")

    # If no OCR engine available or no text produced
    if not extracted_text.strip():
        # Check if the image contains EXIF or PNG text chunks with business metadata
        try:
            from PIL import Image
            img = Image.open(io.BytesIO(content))
            info = img.info
            meta_parts = []
            for k, v in info.items():
                if isinstance(v, str) and len(v) > 5:
                    meta_parts.append(f"{k}: {v}")
            if meta_parts:
                extracted_text = "\n".join(meta_parts)
        except Exception:
            pass

    if not extracted_text.strip():
        return [], [], "This image does not contain readable business text or tabular numbers. Product photos, icons, and non-document images cannot be converted into business data.", ""

    # Parse extracted OCR text
    headers, rows = _parse_text_lines_to_table(extracted_text)

    # Validate if the extracted text genuinely represents business records
    business_keywords = {"price", "cost", "total", "subtotal", "sale", "order", "item", "qty", "quantity", "sku", "rs", "inr", "$", "size", "denim", "dress", "shirt", "top", "jacket"}
    has_business_intent = any(kw in extracted_text.lower() for kw in business_keywords)

    if not rows and not has_business_intent:
        return [], [], "Text was detected in this image, but it does not contain business transactions, prices, or product records.", extracted_text

    warning = None
    if not rows and has_business_intent:
        warning = f"Extracted business text via {ocr_used or 'OCR'}, but structured columns could not be fully resolved."

    return headers, rows, warning, extracted_text


def _parse_text_lines_to_table(text: str) -> Tuple[List[str], List[Dict[str, Any]]]:
    """
    Parses unstructured multi-line text (from PDF, Word, TXT, or OCR)
    looking for tabular structures, delimiter-separated rows, or receipt items.
    """
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return [], []

    # Check for delimiter in lines (e.g. '|', '\t', ',')
    for delim in ["\t", "|", ","]:
        count_delim = [l.count(delim) for l in lines[:10]]
        if count_delim and min(count_delim) >= 2 and max(count_delim) == min(count_delim):
            headers = [c.strip() or f"Column_{i+1}" for i, c in enumerate(lines[0].split(delim))]
            rows = []
            for l in lines[1:]:
                parts = [c.strip() for c in l.split(delim)]
                if not any(parts):
                    continue
                row_dict = {headers[i]: (parts[i] if i < len(parts) else "") for i in range(len(headers))}
                rows.append(row_dict)
            if rows:
                return headers, rows

    # Check for multi-space columnar alignment (common in PDF and receipts)
    sample_split = [re.split(r'\s{2,}', l) for l in lines[:8]]
    col_counts = [len(s) for s in sample_split if len(s) >= 2]
    if len(col_counts) >= 2 and max(col_counts) == min(col_counts):
        col_n = col_counts[0]
        headers = [f"Column_{i+1}" for i in range(col_n)]
        # Check if first line is headers
        first_row = sample_split[0]
        if all(re.match(r'^[A-Za-z\s_-]+$', col) for col in first_row):
            headers = [h.strip() or f"Column_{i+1}" for i, h in enumerate(first_row)]
            data_splits = sample_split[1:]
        else:
            data_splits = sample_split

        rows = []
        for parts in data_splits:
            if not any(parts):
                continue
            row_dict = {headers[i]: (parts[i] if i < len(parts) else "") for i in range(len(headers))}
            rows.append(row_dict)
        if rows:
            return headers, rows

    # Check for line-item receipt pattern: "Item Name ... Price" or "Qty Item Price"
    # e.g., "Vintage Denim Jacket  1200" or "1  Black Oversized Tee  850"
    receipt_pattern = re.compile(r'^(?:(\d+)\s+)?([A-Za-z0-9\s\-_#]+?)\s+([$₹€£]?\s*[\d,]+(?:\.\d{1,2})?)$')
    receipt_rows = []
    ignore_items = {"page", "total", "subtotal", "tax", "vat", "gst", "amount", "balance", "net", "change"}
    for line in lines:
        line_clean = line.strip()
        if re.match(r'^page\s+\d+$', line_clean.lower()):
            continue
        match = receipt_pattern.match(line_clean)
        if match:
            qty, item, price = match.groups()
            cleaned_item = item.strip()
            if cleaned_item.lower() in ignore_items or cleaned_item.isdigit():
                continue
            receipt_rows.append({
                "product_name": cleaned_item,
                "quantity": qty.strip() if qty else "1",
                "retail_price": price.strip()
            })

    if len(receipt_rows) >= 1:
        return ["product_name", "quantity", "retail_price"], receipt_rows

    return [], []


def detect_record_type(columns: List[str], rows: List[Dict[str, Any]]) -> Tuple[str, float]:
    """
    Intelligently determines the primary type of business record:
    - sales_orders
    - products_inventory
    - customer_inquiries
    - social_content
    - mixed_business
    """
    if not columns or not rows:
        return "unknown", 0.0

    cols_lower = [c.lower() for c in columns]
    all_keys_joined = " ".join(cols_lower)

    sales_score = 0
    product_score = 0
    inquiry_score = 0
    social_score = 0

    if any(k in all_keys_joined for k in ["order", "invoice", "sale_date", "sold", "buyer", "customer", "channel"]):
        sales_score += 40
    if any(k in all_keys_joined for k in ["product", "item", "sku", "category", "size", "condition", "brand", "cogs", "cost"]):
        product_score += 40
    if any(k in all_keys_joined for k in ["price", "amount", "revenue", "retail"]):
        sales_score += 20
        product_score += 20
    if any(k in all_keys_joined for k in ["inquiry", "outcome", "enquiry", "status", "notes"]):
        inquiry_score += 45
    if any(k in all_keys_joined for k in ["views", "saves", "likes", "shares", "reel", "caption", "hashtags"]):
        social_score += 50

    total = sales_score + product_score + inquiry_score + social_score
    if total == 0:
        return "general_records", 0.50

    if sales_score > 30 and product_score > 30:
        return "mixed_business", min(0.96, (sales_score + product_score) / 100.0)
    if sales_score >= product_score and sales_score >= inquiry_score and sales_score >= social_score:
        return "sales_orders", min(0.95, sales_score / 60.0)
    if product_score >= sales_score and product_score >= inquiry_score and product_score >= social_score:
        return "products_inventory", min(0.95, product_score / 60.0)
    if social_score > 30:
        return "social_content", min(0.95, social_score / 50.0)
    if inquiry_score > 30:
        return "customer_inquiries", min(0.95, inquiry_score / 45.0)

    return "mixed_business", 0.70
